"""§496/§502/§788 导入过程:XLSX 只读检查与受控快照取数。

- 未指定工作表:列出全部工作表(名称、行数、列数),由用户指定分析对象;
- 指定工作表:输出表头列名与数据行(供上层物化受控 CSV 快照);
- 公式单元格缺少已缓存结果时计入 formula_unavailable,快照中为空值,
  绝不替产品宣称会自动计算 Excel 公式;
- 只读打开,绝不修改输入文件;缺失 openpyxl 时如实报告(环境准备项,
  本命令绝不安装依赖);行数超出 max_rows 时按 row_limit 如实上报。
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime, time
from typing import Any


def _cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return value


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def inspect_xlsx(path: str, sheet: str | None, max_rows: int) -> dict[str, Any]:
    import openpyxl  # 延迟导入:缺失时走环境准备项分支

    if sheet is None:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = [
                {"name": name, "rows": workbook[name].max_row or 0, "cols": workbook[name].max_column or 0}
                for name in workbook.sheetnames
            ]
        finally:
            workbook.close()
        return {"ok": True, "format": "xlsx", "sheets": sheets}

    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet not in values.sheetnames:
            return {
                "ok": False,
                "error_kind": "sheet_not_found",
                "reason": f"工作表不存在: {sheet}",
                "sheets": values.sheetnames,
            }
        ws_values = values[sheet]
        ws_formulas = formulas[sheet]

        header: list[str] | None = None
        rows: list[list[Any]] = []
        formula_unavailable = 0
        row_count = 0
        value_rows = ws_values.iter_rows(values_only=True)
        formula_rows = ws_formulas.iter_rows(values_only=True)
        for value_row, formula_row in zip(value_rows, formula_rows):
            if header is None:
                header = [
                    str(cell).strip() if cell is not None and str(cell).strip() != "" else f"column_{index + 1}"
                    for index, cell in enumerate(value_row)
                ]
                continue
            row_count += 1
            if row_count > max_rows:
                return {
                    "ok": False,
                    "error_kind": "row_limit",
                    "reason": f"数据行数超出导入上限: {max_rows}",
                }
            record: list[Any] = []
            for index, cell in enumerate(value_row):
                formula_cell = formula_row[index] if index < len(formula_row) else None
                if cell is None and _is_formula(formula_cell):
                    formula_unavailable += 1  # §496:公式值缺少已缓存结果 → unavailable
                record.append(_cell_value(cell))
            rows.append(record)
    finally:
        values.close()
        formulas.close()

    return {
        "ok": True,
        "format": "xlsx",
        "sheet": sheet,
        "row_count": row_count,
        "columns": header or [],
        "rows": rows,
        "formula_unavailable": formula_unavailable,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="analyst_runtime inspect-xlsx",
        description="只读解析 XLSX 工作表清单或指定工作表数据,绝不修改文件",
    )
    parser.add_argument("--input", required=True, help="XLSX 文件路径")
    parser.add_argument("--sheet", default=None, help="工作表名;缺省时仅列出工作表")
    parser.add_argument("--max-rows", type=int, default=100_000, help="数据行数上限")
    args = parser.parse_args(argv)

    try:
        report = inspect_xlsx(args.input, args.sheet, args.max_rows)
    except ImportError:
        report = {
            "ok": False,
            "error_kind": "environment",
            "reason": "openpyxl 未安装(环境准备项,本命令不安装依赖)",
        }
    except Exception as error:  # noqa: BLE001 — 统一转为结构化失败输出
        report = {"ok": False, "error_kind": "parse", "reason": f"解析失败: {error}"}

    json.dump(report, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
